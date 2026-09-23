"""Phase 21A — XLSX security policy and container-validation foundation."""

from __future__ import annotations

import io
import zipfile
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app.domain.attachment_policy import (
    OPENPYXL_DATA_ONLY,
    OPENPYXL_KEEP_VBA,
    OPENPYXL_READ_ONLY,
    XLSX_MAX_ZIP_ENTRIES,
    detect_attachment_kind,
    evaluate_attachment_content,
    evaluate_attachment_metadata,
    validate_xlsx_container,
    xlsx_product_analysis_enabled,
)
from app.domain.enums import AttachmentKind
from app.domain.exceptions import (
    AttachmentExceedsLimitError,
    AttachmentUnsupportedError,
)
from app.domain.models import AttachmentContent, AttachmentMetadata
from tests.unit.infrastructure.attachments.fixtures import (
    minimal_xlsx_bytes,
    xlsx_with_external_relationship,
    xlsx_with_formula_cells,
)

_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _metadata(
    *,
    filename: str = "budget.xlsx",
    media_type: str = _XLSX_MEDIA,
    reported_size: int = 32,
) -> AttachmentMetadata:
    return AttachmentMetadata(
        provider_attachment_id="att-xlsx",
        filename=filename,
        media_type=media_type,
        reported_size=reported_size,
    )


def _content(payload: bytes, metadata: AttachmentMetadata | None = None) -> AttachmentContent:
    item = metadata or _metadata(reported_size=len(payload))
    return AttachmentContent(
        metadata=item,
        content=payload,
        source_message_id="msg-1",
        source_attachment_id=item.provider_attachment_id,
    )


def _zip_with_entries(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, mode="w", compression=zipfile.ZIP_STORED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def test_product_analysis_remains_disabled_for_xlsx_when_gate_disabled(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Retain the A/B fail-closed gate guarantee; Phase 21D enables it by default.
    monkeypatch.setattr("app.domain.attachment_policy._PRODUCT_ANALYSIS_KINDS", frozenset())
    assert xlsx_product_analysis_enabled() is False
    payload = minimal_xlsx_bytes()
    with pytest.raises(AttachmentUnsupportedError):
        evaluate_attachment_metadata(
            _metadata(reported_size=len(payload)),
        )


def test_valid_minimal_xlsx_container_passes_validation() -> None:
    payload = minimal_xlsx_bytes()
    result = validate_xlsx_container(payload)
    assert result.recognized is True
    assert result.compressed_size == len(payload)
    assert result.archive_member_count >= 2
    assert result.aggregate_uncompressed_size > 0
    assert result.warnings == ()


def test_openpyxl_workbook_container_passes_and_detects_kind() -> None:
    payload = xlsx_with_formula_cells()
    result = validate_xlsx_container(payload)
    assert result.recognized is True
    assert detect_attachment_kind(payload) is AttachmentKind.XLSX
    metadata = _metadata(reported_size=len(payload))
    assert evaluate_attachment_content(_content(payload, metadata)) is AttachmentKind.XLSX


def test_xlsx_extension_with_non_zip_content_fails() -> None:
    payload = b"not-a-zip-workbook"
    metadata = _metadata(reported_size=len(payload))
    with pytest.raises(AttachmentUnsupportedError):
        evaluate_attachment_content(_content(payload, metadata))
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_arbitrary_zip_renamed_xlsx_fails() -> None:
    payload = _zip_with_entries({"readme.txt": b"hello"})
    metadata = _metadata(reported_size=len(payload))
    with pytest.raises(AttachmentUnsupportedError):
        evaluate_attachment_content(_content(payload, metadata))
    assert detect_attachment_kind(payload) is None


def test_malformed_zip_fails() -> None:
    payload = b"PK\x03\x04" + b"\x00" * 32
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_mime_mismatch_fails_metadata() -> None:
    with pytest.raises(AttachmentUnsupportedError):
        evaluate_attachment_metadata(
            _metadata(filename="budget.xlsx", media_type="application/pdf", reported_size=100)
        )
    with pytest.raises(AttachmentUnsupportedError):
        evaluate_attachment_metadata(
            _metadata(
                filename="budget.xlsx",
                media_type="application/octet-stream",
                reported_size=100,
            )
        )


@pytest.mark.parametrize(
    ("filename", "media_type"),
    [
        (".xls", "application/vnd.ms-excel"),
        ("legacy.xls", "application/vnd.ms-excel"),
        (
            "macro.xlsm",
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        ),
        (
            "binary.xlsb",
            "application/vnd.ms-excel.sheet.binary.macroEnabled.12",
        ),
        ("rows.csv", "text/csv"),
        ("rows.tsv", "text/tab-separated-values"),
    ],
)
def test_unsupported_spreadsheet_formats_fail_closed(filename: str, media_type: str) -> None:
    with pytest.raises(AttachmentUnsupportedError):
        evaluate_attachment_metadata(
            _metadata(filename=filename, media_type=media_type, reported_size=128)
        )


def test_docx_and_xlsx_are_discriminated() -> None:
    xlsx_payload = minimal_xlsx_bytes()
    docx_payload = _zip_with_entries(
        {
            "[Content_Types].xml": b"<Types></Types>",
            "word/document.xml": b"<w:document></w:document>",
        }
    )
    assert detect_attachment_kind(xlsx_payload) is AttachmentKind.XLSX
    assert detect_attachment_kind(docx_payload) is AttachmentKind.DOCX
    polyglot = _zip_with_entries(
        {
            "[Content_Types].xml": b"<Types></Types>",
            "word/document.xml": b"<w:document></w:document>",
            "xl/workbook.xml": b"<workbook></workbook>",
        }
    )
    assert detect_attachment_kind(polyglot) is None


def test_excessive_archive_member_count_fails() -> None:
    entries = {
        "[Content_Types].xml": b"<Types></Types>",
        "xl/workbook.xml": b"<workbook></workbook>",
    }
    for index in range(XLSX_MAX_ZIP_ENTRIES):
        entries[f"xl/pad/part{index}.xml"] = b"x"
    payload = _zip_with_entries(entries)
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_excessive_aggregate_uncompressed_size_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "app.domain.attachment_policy.XLSX_MAX_UNCOMPRESSED_TOTAL",
        800,
    )
    payload = minimal_xlsx_bytes(
        extra={
            "xl/a.xml": b"a" * 500,
            "xl/b.xml": b"b" * 500,
        }
    )
    with pytest.raises(AttachmentExceedsLimitError):
        validate_xlsx_container(payload)


def test_excessive_single_part_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.domain.attachment_policy.XLSX_MAX_SINGLE_PART", 100)
    payload = minimal_xlsx_bytes(extra={"xl/huge.xml": b"h" * 200})
    with pytest.raises(AttachmentExceedsLimitError):
        validate_xlsx_container(payload)


def test_excessive_compression_ratio_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.domain.attachment_policy.XLSX_MAX_COMPRESSION_RATIO", 2)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types></Types>")
        archive.writestr("xl/workbook.xml", b"<workbook></workbook>")
        archive.writestr("xl/bomb.xml", b"\x00" * 10_000)
    payload = buffer.getvalue()
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_path_traversal_member_name_fails() -> None:
    payload = _zip_with_entries(
        {
            "[Content_Types].xml": b"<Types></Types>",
            "xl/workbook.xml": b"<workbook></workbook>",
            "../evil.txt": b"nope",
        }
    )
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_absolute_member_path_fails() -> None:
    payload = _zip_with_entries(
        {
            "[Content_Types].xml": b"<Types></Types>",
            "xl/workbook.xml": b"<workbook></workbook>",
            "/abs/evil.txt": b"nope",
        }
    )
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_vba_project_rejected() -> None:
    payload = minimal_xlsx_bytes(include_vba=True)
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_encryption_marker_rejected() -> None:
    payload = minimal_xlsx_bytes(extra={"EncryptionInfo": b"secret"})
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_external_relationship_rejected() -> None:
    payload = xlsx_with_external_relationship()
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_external_links_directory_rejected() -> None:
    payload = minimal_xlsx_bytes(
        extra={"xl/externalLinks/externalLink1.xml": b"<externalLink/>"}
    )
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_data_connections_rejected() -> None:
    payload = minimal_xlsx_bytes(extra={"xl/connections.xml": b"<connections/>"})
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)


def test_embedded_object_warns_but_allows_container() -> None:
    payload = minimal_xlsx_bytes(extra={"xl/embeddings/oleObject1.bin": b"ole"})
    result = validate_xlsx_container(payload)
    assert result.recognized is True
    assert "xlsx_embedded_object_ignored" in result.warnings


def test_macro_content_types_marker_rejected() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, mode="w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            (
                b'<?xml version="1.0"?><Types>'
                b'<Override PartName="/xl/workbook.xml" '
                b'ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml"/>'
                b"</Types>"
            ),
        )
        archive.writestr("xl/workbook.xml", b"<workbook></workbook>")
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(buffer.getvalue())


def test_formulas_are_not_executed_by_openpyxl_safe_settings() -> None:
    payload = xlsx_with_formula_cells()
    # Container validation must not require openpyxl and must not calculate.
    validate_xlsx_container(payload)
    workbook = load_workbook(
        io.BytesIO(payload),
        read_only=OPENPYXL_READ_ONLY,
        data_only=OPENPYXL_DATA_ONLY,
        keep_vba=OPENPYXL_KEEP_VBA,
    )
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(min_row=1, max_row=3, max_col=2, values_only=True))
        assert rows[2][0] == "=A1+A2"
        assert isinstance(rows[0][1], str)
        assert rows[0][1].startswith("=HYPERLINK")
        # Cached computed value 3 must not appear when data_only=False.
        assert rows[2][0] != 3
    finally:
        workbook.close()


def test_validation_performs_no_network_call(monkeypatch: pytest.MonkeyPatch) -> None:
    blocked = MagicMock(side_effect=AssertionError("network forbidden"))
    monkeypatch.setattr("socket.socket", blocked)
    payload = xlsx_with_external_relationship()
    with pytest.raises(AttachmentUnsupportedError):
        validate_xlsx_container(payload)
    payload_ok = minimal_xlsx_bytes()
    validate_xlsx_container(payload_ok)
    blocked.assert_not_called()


def test_xlsx_content_validation_does_not_extract_cells() -> None:
    payload = xlsx_with_formula_cells()
    metadata = _metadata(reported_size=len(payload))
    kind = evaluate_attachment_content(_content(payload, metadata))
    assert kind is AttachmentKind.XLSX
    # No ParsedAttachment / cell dump is produced by Phase 21A APIs.
    result = validate_xlsx_container(payload)
    assert not hasattr(result, "cells")
    assert result.warnings == ()
