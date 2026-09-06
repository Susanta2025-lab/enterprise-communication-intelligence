"""Synthetic attachment fixtures. No real documents or malware samples."""

from __future__ import annotations

import io
import struct
import zipfile
import zlib

from docx import Document
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.domain.enums import AttachmentDisposition
from app.domain.models import AttachmentContent, AttachmentMetadata


def attachment_content(
    payload: bytes,
    *,
    filename: str,
    media_type: str,
    attachment_id: str = "att-1",
    message_id: str = "fake-msg-001",
) -> AttachmentContent:
    return AttachmentContent(
        metadata=AttachmentMetadata(
            provider_attachment_id=attachment_id,
            filename=filename,
            media_type=media_type,
            reported_size=len(payload),
            disposition=AttachmentDisposition.ATTACHMENT,
        ),
        content=payload,
        source_message_id=message_id,
        source_attachment_id=attachment_id,
    )


def pdf_with_text(*page_texts: str) -> bytes:
    writer = PdfWriter()
    for text in page_texts:
        page = writer.add_blank_page(width=300, height=300)
        font = writer._add_object(
            DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 36 250 Td ({escaped}) Tj ET".encode("latin-1", "replace"))
        page[NameObject("/Contents")] = writer._add_object(content)
    return _write_pdf(writer)


def pdf_blank_pages(count: int) -> bytes:
    writer = PdfWriter()
    for _ in range(count):
        writer.add_blank_page(width=200, height=200)
    return _write_pdf(writer)


def pdf_encrypted() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    return _write_pdf(writer)


def pdf_with_ignored_actions(text: str = "Visible contract text") -> bytes:
    payload = pdf_with_text(text)
    header_end = payload.find(b"\n")
    comment = b"% /JS (app.alert('x')) /JavaScript /Launch /EmbeddedFile /URI\n"
    if header_end == -1:
        return comment + payload
    return payload[: header_end + 1] + comment + payload[header_end + 1 :]


def docx_with_text(*paragraphs: str, tables: list[list[list[str]]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for table_data in tables or []:
        table = document.add_table(rows=len(table_data), cols=len(table_data[0]))
        for row_index, row in enumerate(table_data):
            for col_index, value in enumerate(row):
                table.rows[row_index].cells[col_index].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_with_zip_member(name: str, member_bytes: bytes = b"ignored") -> bytes:
    payload = bytearray(docx_with_text("Visible paragraph"))
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload), "r") as source:
        with zipfile.ZipFile(buffer, "w") as dest:
            for info in source.infolist():
                dest.writestr(info, source.read(info.filename))
            dest.writestr(name, member_bytes)
    return buffer.getvalue()


def docx_with_external_relationship() -> bytes:
    payload = docx_with_text("Contract body")
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload), "r") as source:
        with zipfile.ZipFile(buffer, "w") as dest:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename.endswith("document.xml.rels"):
                    external = (
                        b'<Relationship Id="rEvil" '
                        b'Type="http://schemas.openxmlformats.org/officeDocument/'
                        b'2006/relationships/hyperlink" '
                        b'Target="https://evil.example/ignore" '
                        b'TargetMode="External"/>'
                    )
                    data = data.replace(b"</Relationships>", external + b"</Relationships>")
                dest.writestr(info, data)
    return buffer.getvalue()


def tiny_jpeg() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (16, 16), color=(12, 34, 56)).save(buffer, format="JPEG")
    return buffer.getvalue()


def tiny_png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (16, 16), color=(9, 8, 7)).save(buffer, format="PNG")
    return buffer.getvalue()


def jpeg_with_exif_secret(secret: bytes = b"ECI-EXIF-GPS-SECRET") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (16, 16), color=(1, 2, 3)).save(buffer, format="JPEG")
    data = buffer.getvalue()
    app1 = b"\xff\xe1" + struct.pack(">H", len(secret) + 8) + b"Exif\x00\x00" + secret
    return data[:2] + app1 + data[2:]


def png_with_dimensions(width: int, height: int) -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        crc = zlib.crc32(tag + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", crc)

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\x00\x00\x00")
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def _write_pdf(writer: PdfWriter) -> bytes:
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
