"""Bounded text extraction for already-validated DOCX attachments."""

from __future__ import annotations

import io
import zipfile

from defusedxml import ElementTree
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.domain.attachment_policy import MAX_EXTRACTED_TEXT_CHARS
from app.domain.enums import AttachmentKind
from app.domain.exceptions import (
    AttachmentNoExtractableTextError,
    AttachmentParseError,
    AttachmentUnsupportedError,
)
from app.domain.models.attachment import AttachmentContent, ParsedAttachment

_EMBEDDING_PREFIXES = ("word/embeddings/", "word/oleobjects/")


def parse_docx_attachment(content: AttachmentContent) -> ParsedAttachment:
    """Extract paragraph and table text. Macros and embeddings are not opened."""
    payload = content.content
    warnings = _docx_structure_warnings(payload)
    try:
        document = Document(io.BytesIO(payload))
    except (PackageNotFoundError, ValueError, OSError, KeyError) as exc:
        raise AttachmentParseError() from exc
    except Exception as exc:
        if isinstance(exc, (AttachmentParseError, AttachmentUnsupportedError)):
            raise
        raise AttachmentParseError() from exc

    collected: list[str] = []
    used = 0
    truncated = False
    try:
        for paragraph in document.paragraphs:
            piece = (paragraph.text or "").strip()
            if not piece:
                continue
            remaining = MAX_EXTRACTED_TEXT_CHARS - used
            if remaining <= 0:
                truncated = True
                break
            if len(piece) > remaining:
                collected.append(piece[:remaining])
                used += remaining
                truncated = True
                break
            collected.append(piece)
            used += len(piece)
        if not truncated:
            for table in document.tables:
                for row in table.rows:
                    cells = [(cell.text or "").strip() for cell in row.cells]
                    piece = "\t".join(cell for cell in cells if cell)
                    if not piece:
                        continue
                    remaining = MAX_EXTRACTED_TEXT_CHARS - used
                    if remaining <= 0:
                        truncated = True
                        break
                    if len(piece) > remaining:
                        collected.append(piece[:remaining])
                        used += remaining
                        truncated = True
                        break
                    collected.append(piece)
                    used += len(piece)
                if truncated:
                    break
    except Exception as exc:
        raise AttachmentParseError() from exc

    text = "\n".join(collected).strip()
    if not text:
        raise AttachmentNoExtractableTextError()

    return ParsedAttachment(
        kind=AttachmentKind.DOCX,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extracted_text=text,
        page_count=None,
        character_count=len(text),
        truncated=truncated,
        warnings=warnings,
    )


def _docx_structure_warnings(payload: bytes) -> tuple[str, ...]:
    warnings: list[str] = []
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload))
    except zipfile.BadZipFile as exc:
        raise AttachmentParseError() from exc
    with archive:
        names = {info.filename.replace("\\", "/") for info in archive.infolist()}
        if any(
            name.startswith(_EMBEDDING_PREFIXES) or "/embeddings/" in name
            for name in names
        ):
            warnings.append("docx_embedded_object_ignored")
        for name in names:
            lowered = name.lower()
            if not lowered.endswith(".rels") and "document.xml.rels" not in lowered:
                continue
            try:
                rels = archive.read(name)
                root = ElementTree.fromstring(rels)
            except Exception:
                continue
            if any(element.get("TargetMode") == "External" for element in root.iter()):
                warnings.append("docx_external_relationship_ignored")
                break
    return tuple(warnings)
