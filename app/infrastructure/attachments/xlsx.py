"""Bounded XLSX workbook extraction for already-validated containers.

Phase 21B: validate_xlsx_container → openpyxl (read-only, no formula execution)
→ bounded sheet/row/column/cell/character sample → ParsedAttachment.

Formulas, hyperlinks, and external references are never followed or evaluated.
Comments/notes are out of Phase 21B scope under ``read_only=True`` (ignored
deterministically). Cell hyperlink objects are unavailable in read-only mode;
``HYPERLINK(...)`` formulas are retained as inert formula text only.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell
from openpyxl.utils.exceptions import InvalidFileException

from app.domain.attachment_policy import (
    MAX_EXTRACTED_TEXT_CHARS,
    OPENPYXL_DATA_ONLY,
    OPENPYXL_KEEP_VBA,
    OPENPYXL_READ_ONLY,
    XLSX_MAX_CELL_CHARS,
    XLSX_MAX_COLUMNS_PER_SHEET,
    XLSX_MAX_DATA_ROWS_PER_SHEET,
    XLSX_MAX_SHEETS,
    XLSX_MAX_TOTAL_CELLS,
    validate_xlsx_container,
)
from app.domain.enums import AttachmentKind
from app.domain.exceptions import (
    AttachmentNoExtractableTextError,
    AttachmentParseError,
    AttachmentUnsupportedError,
)
from app.domain.models.attachment import AttachmentContent, ParsedAttachment
from app.domain.models.workbook_extraction import (
    SheetExtraction,
    WorkbookExtraction,
    XlsxCellValueType,
    XlsxSheetVisibility,
    XlsxTruncationReason,
)

_XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
_HYPERLINK_FORMULA_RE = re.compile(r"^\s*=\s*HYPERLINK\s*\(", re.IGNORECASE)
_EXTERNAL_REF_RE = re.compile(r"\[[^\]]+\]")
_ROW_JOIN = " | "


def parse_xlsx_attachment(content: AttachmentContent) -> ParsedAttachment:
    """Extract a bounded tabular sample from a Phase 21A-validated XLSX payload."""
    extraction = extract_xlsx_workbook(content)
    if extraction.cells_emitted <= 0:
        raise AttachmentNoExtractableTextError()
    text = serialize_workbook_extraction(extraction)
    if not text.strip():
        raise AttachmentNoExtractableTextError()

    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        text = text[:MAX_EXTRACTED_TEXT_CHARS]
        extraction = extraction.model_copy(
            update={
                "truncated": True,
                "truncation_reasons": _append_unique(
                    extraction.truncation_reasons,
                    XlsxTruncationReason.CHARACTERS.value,
                ),
                "warnings": _append_unique(
                    extraction.warnings,
                    XlsxTruncationReason.CHARACTERS.value,
                ),
                "characters_emitted": MAX_EXTRACTED_TEXT_CHARS,
            }
        )

    return ParsedAttachment(
        kind=AttachmentKind.XLSX,
        media_type=_XLSX_MEDIA_TYPE,
        extracted_text=text,
        page_count=extraction.sheet_count,
        character_count=len(text),
        truncated=extraction.truncated,
        warnings=extraction.warnings,
    )


def extract_xlsx_workbook(content: AttachmentContent) -> WorkbookExtraction:
    """Return the structured workbook sample without AI or persistence.

    Requires a Phase 21A-valid container. Suitable for unit tests and Phase 21C.
    """
    payload = content.content
    try:
        container = validate_xlsx_container(payload)
    except (AttachmentUnsupportedError, AttachmentParseError):
        raise
    except Exception as exc:
        raise AttachmentParseError() from exc

    warnings: list[str] = list(container.warnings)
    try:
        workbook = load_workbook(
            io.BytesIO(payload),
            read_only=OPENPYXL_READ_ONLY,
            data_only=OPENPYXL_DATA_ONLY,
            keep_vba=OPENPYXL_KEEP_VBA,
            keep_links=False,
        )
    except (InvalidFileException, OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise AttachmentParseError() from exc
    except Exception as exc:
        if isinstance(
            exc,
            (
                AttachmentParseError,
                AttachmentUnsupportedError,
                AttachmentNoExtractableTextError,
            ),
        ):
            raise
        raise AttachmentParseError() from exc

    try:
        try:
            return _extract_workbook(workbook, warnings)
        except Exception as exc:
            if isinstance(
                exc,
                (
                    AttachmentParseError,
                    AttachmentUnsupportedError,
                    AttachmentNoExtractableTextError,
                ),
            ):
                raise
            raise AttachmentParseError() from exc
    finally:
        workbook.close()


def serialize_workbook_extraction(extraction: WorkbookExtraction) -> str:
    """Deterministic plain-text serialization for ``attachment_texts``."""
    lines: list[str] = [
        "WORKBOOK",
        f"kind: {extraction.kind}",
        f"sheet_count: {extraction.sheet_count}",
        f"sheets_processed: {extraction.sheets_processed}",
        f"sheets_skipped: {extraction.sheets_skipped}",
        f"truncated: {'true' if extraction.truncated else 'false'}",
        f"cells_emitted: {extraction.cells_emitted}",
        f"warnings: {_format_warning_list(extraction.warnings)}",
        f"truncation_reasons: {_format_warning_list(extraction.truncation_reasons)}",
        "",
    ]
    sheet_number = 0
    for sheet in extraction.sheets:
        if sheet.skipped:
            lines.append(f"SHEET SKIPPED index={sheet.index}")
            lines.append(f"name: {sheet.name}")
            lines.append(f"visibility: {sheet.visibility.value}")
            lines.append(f"skip_reason: {sheet.skip_reason or ''}")
            lines.append("")
            continue
        sheet_number += 1
        lines.append(f"SHEET {sheet_number}")
        lines.append(f"name: {sheet.name}")
        lines.append(f"visibility: {sheet.visibility.value}")
        if sheet.source_max_row is not None or sheet.source_max_column is not None:
            row_part = (
                str(sheet.source_max_row) if sheet.source_max_row is not None else ""
            )
            col_part = (
                str(sheet.source_max_column)
                if sheet.source_max_column is not None
                else ""
            )
            lines.append(f"dimensions: rows={row_part} cols={col_part}")
        lines.append(f"header_inferred: {'true' if sheet.header_inferred else 'false'}")
        if sheet.header_inferred:
            lines.append(f"header: {_join_row(sheet.header_cells)}")
        lines.append(f"sampled_data_rows: {len(sheet.data_rows)}")
        lines.append(f"columns: {sheet.columns_processed}")
        lines.append(f"truncated: {'true' if sheet.truncated else 'false'}")
        if sheet.truncation_reasons:
            lines.append(
                f"truncation_reasons: {_format_warning_list(sheet.truncation_reasons)}"
            )
        lines.append("---")
        for row in sheet.data_rows:
            lines.append(_join_row(row))
        lines.append("")
    return "\n".join(lines).rstrip() + ("\n" if lines else "")


def _extract_workbook(workbook: Any, seed_warnings: list[str]) -> WorkbookExtraction:
    warnings = list(seed_warnings)
    truncation_reasons: list[str] = []
    sheets_out: list[SheetExtraction] = []
    sheet_names: list[str] = list(workbook.sheetnames)
    sheet_count = len(sheet_names)
    cells_emitted = 0
    characters_emitted = 0
    truncated = False
    stop_workbook = False
    sheets_processed = 0
    sheets_skipped = 0

    if sheet_count > XLSX_MAX_SHEETS:
        truncated = True
        _add_reason(truncation_reasons, warnings, XlsxTruncationReason.SHEETS.value)

    for index, name in enumerate(sheet_names):
        if stop_workbook:
            sheets_skipped += 1
            sheets_out.append(
                SheetExtraction(
                    name=name,
                    visibility=_sheet_visibility(workbook[name]),
                    index=index,
                    skipped=True,
                    skip_reason=XlsxTruncationReason.CELLS.value,
                )
            )
            continue
        if index >= XLSX_MAX_SHEETS:
            sheets_skipped += 1
            sheets_out.append(
                SheetExtraction(
                    name=name,
                    visibility=_sheet_visibility(workbook[name]),
                    index=index,
                    skipped=True,
                    skip_reason=XlsxTruncationReason.SHEETS.value,
                )
            )
            continue

        worksheet = workbook[name]
        visibility = _sheet_visibility(worksheet)
        if visibility is not XlsxSheetVisibility.VISIBLE:
            _add_warning(warnings, "xlsx_hidden_sheets_present")

        sheet, cells_emitted, characters_emitted, stop_workbook = _extract_sheet(
            worksheet,
            name=name,
            index=index,
            visibility=visibility,
            cells_emitted=cells_emitted,
            characters_emitted=characters_emitted,
            warnings=warnings,
        )
        if sheet.truncated:
            truncated = True
            for reason in sheet.truncation_reasons:
                _add_reason(truncation_reasons, warnings, reason)
        sheets_out.append(sheet)
        sheets_processed += 1

    if stop_workbook:
        truncated = True
        _add_reason(truncation_reasons, warnings, XlsxTruncationReason.CELLS.value)

    return WorkbookExtraction(
        sheet_count=sheet_count,
        sheets_processed=sheets_processed,
        sheets_skipped=sheets_skipped,
        sheets=tuple(sheets_out),
        warnings=tuple(warnings),
        truncated=truncated or bool(truncation_reasons),
        truncation_reasons=tuple(truncation_reasons),
        cells_emitted=cells_emitted,
        characters_emitted=characters_emitted,
    )


def _extract_sheet(
    worksheet: Any,
    *,
    name: str,
    index: int,
    visibility: XlsxSheetVisibility,
    cells_emitted: int,
    characters_emitted: int,
    warnings: list[str],
) -> tuple[SheetExtraction, int, int, bool]:
    truncation_reasons: list[str] = []
    header_cells: tuple[str, ...] = ()
    header_inferred = False
    data_rows: list[tuple[str, ...]] = []
    columns_processed = 0
    sheet_cells = 0
    stop_workbook = False
    truncated = False

    source_max_row = _safe_int_attr(worksheet, "max_row")
    source_max_column = _safe_int_attr(worksheet, "max_column")

    # Never allocate from declared dimensions; only use them as truncation hints.
    if (
        source_max_column is not None
        and source_max_column > XLSX_MAX_COLUMNS_PER_SHEET
    ):
        truncated = True
        _add_reason(truncation_reasons, warnings, XlsxTruncationReason.COLUMNS.value)
    scan_rows = XLSX_MAX_DATA_ROWS_PER_SHEET + 1  # header + data
    if source_max_row is not None and source_max_row > scan_rows:
        truncated = True
        _add_reason(truncation_reasons, warnings, XlsxTruncationReason.ROWS.value)

    try:
        row_iter = worksheet.iter_rows(
            min_row=1,
            max_row=scan_rows,
            max_col=XLSX_MAX_COLUMNS_PER_SHEET,
            values_only=False,
        )
    except Exception as exc:
        raise AttachmentParseError() from exc

    for row in row_iter:
        if stop_workbook:
            break
        if cells_emitted >= XLSX_MAX_TOTAL_CELLS:
            truncated = True
            stop_workbook = True
            _add_reason(truncation_reasons, warnings, XlsxTruncationReason.CELLS.value)
            break
        if characters_emitted >= MAX_EXTRACTED_TEXT_CHARS:
            truncated = True
            stop_workbook = True
            _add_reason(truncation_reasons, warnings, XlsxTruncationReason.CHARACTERS.value)
            break

        rendered: list[str] = []
        row_has_value = False
        row_width = 0
        for col_index, cell in enumerate(row, start=1):
            if cells_emitted >= XLSX_MAX_TOTAL_CELLS:
                truncated = True
                stop_workbook = True
                _add_reason(truncation_reasons, warnings, XlsxTruncationReason.CELLS.value)
                break
            if characters_emitted >= MAX_EXTRACTED_TEXT_CHARS:
                truncated = True
                stop_workbook = True
                _add_reason(
                    truncation_reasons, warnings, XlsxTruncationReason.CHARACTERS.value
                )
                break

            text, value_type, formula_text, cell_truncated = _normalize_cell(cell)
            if cell_truncated:
                truncated = True
                _add_reason(
                    truncation_reasons, warnings, XlsxTruncationReason.CELL_CHARS.value
                )
            if formula_text is not None:
                _add_warning(warnings, "xlsx_formulas_present")
                if _HYPERLINK_FORMULA_RE.match(formula_text):
                    _add_warning(warnings, "xlsx_hyperlink_formula")
                if _EXTERNAL_REF_RE.search(formula_text):
                    _add_warning(warnings, "xlsx_external_reference")

            if value_type is XlsxCellValueType.BLANK or text == "":
                rendered.append("")
                continue

            remaining = MAX_EXTRACTED_TEXT_CHARS - characters_emitted
            if len(text) > remaining:
                text = text[:remaining]
                truncated = True
                stop_workbook = True
                _add_reason(truncation_reasons, warnings, XlsxTruncationReason.CHARACTERS.value)

            row_has_value = True
            row_width = col_index
            rendered.append(text)
            cells_emitted += 1
            sheet_cells += 1
            characters_emitted += len(text)

        if not row_has_value:
            continue

        clipped = rendered[:row_width]
        columns_processed = max(columns_processed, len(clipped))
        row_tuple = tuple(clipped)

        if not header_inferred:
            header_cells = row_tuple
            header_inferred = True
            continue

        if len(data_rows) >= XLSX_MAX_DATA_ROWS_PER_SHEET:
            truncated = True
            _add_reason(truncation_reasons, warnings, XlsxTruncationReason.ROWS.value)
            break

        data_rows.append(row_tuple)

    rows_processed = (1 if header_inferred else 0) + len(data_rows)
    return (
        SheetExtraction(
            name=name,
            visibility=visibility,
            index=index,
            header_inferred=header_inferred,
            header_cells=header_cells,
            data_rows=tuple(data_rows),
            rows_processed=rows_processed,
            columns_processed=columns_processed,
            cells_emitted=sheet_cells,
            truncated=truncated,
            truncation_reasons=tuple(truncation_reasons),
            source_max_row=source_max_row,
            source_max_column=source_max_column,
        ),
        cells_emitted,
        characters_emitted,
        stop_workbook,
    )


def _normalize_cell(
    cell: Any,
) -> tuple[str, XlsxCellValueType, str | None, bool]:
    """Return (text, type, raw_formula_or_none, cell_char_truncated)."""
    if cell is None or isinstance(cell, EmptyCell):
        return "", XlsxCellValueType.BLANK, None, False

    value = getattr(cell, "value", None)
    if value is None:
        return "", XlsxCellValueType.BLANK, None, False

    data_type = getattr(cell, "data_type", None)
    if data_type == "f" or (isinstance(value, str) and value.startswith("=")):
        formula = value if isinstance(value, str) else str(value)
        text, truncated = _bound_cell_text(f"formula:{formula}")
        return text, XlsxCellValueType.FORMULA, formula, truncated

    if data_type == "e" or (
        isinstance(value, str) and value.startswith("#") and value.endswith("!")
    ):
        text, truncated = _bound_cell_text(str(value))
        return text, XlsxCellValueType.ERROR, None, truncated

    if isinstance(value, bool):
        return ("true" if value else "false"), XlsxCellValueType.BOOLEAN, None, False

    if isinstance(value, datetime):
        text = value.replace(microsecond=0).isoformat()
        text, truncated = _bound_cell_text(text)
        return text, XlsxCellValueType.DATETIME, None, truncated

    if isinstance(value, date) and not isinstance(value, datetime):
        text, truncated = _bound_cell_text(value.isoformat())
        return text, XlsxCellValueType.DATE, None, truncated

    if isinstance(value, time):
        text, truncated = _bound_cell_text(value.replace(microsecond=0).isoformat())
        return text, XlsxCellValueType.TIME, None, truncated

    if isinstance(value, int) and not isinstance(value, bool):
        text, truncated = _bound_cell_text(str(value))
        return text, XlsxCellValueType.INTEGER, None, truncated

    if isinstance(value, float):
        text, truncated = _bound_cell_text(_format_float(value))
        return text, XlsxCellValueType.FLOAT, None, truncated

    if isinstance(value, Decimal):
        text, truncated = _bound_cell_text(format(value, "f"))
        return text, XlsxCellValueType.FLOAT, None, truncated

    text, truncated = _bound_cell_text(str(value))
    return text, XlsxCellValueType.STRING, None, truncated


def _bound_cell_text(text: str) -> tuple[str, bool]:
    cleaned = text.replace("\x00", "")
    if len(cleaned) <= XLSX_MAX_CELL_CHARS:
        return cleaned, False
    return cleaned[:XLSX_MAX_CELL_CHARS], True


def _format_float(value: float) -> str:
    if value.is_integer():
        return str(int(value))
    return format(value, "g")


def _sheet_visibility(worksheet: Any) -> XlsxSheetVisibility:
    state = str(getattr(worksheet, "sheet_state", "visible") or "visible").lower()
    if state == "hidden":
        return XlsxSheetVisibility.HIDDEN
    if state in {"veryhidden", "very_hidden"}:
        return XlsxSheetVisibility.VERY_HIDDEN
    return XlsxSheetVisibility.VISIBLE


def _safe_int_attr(obj: Any, name: str) -> int | None:
    value = getattr(obj, name, None)
    if isinstance(value, int) and value >= 0:
        return value
    return None


def _join_row(cells: tuple[str, ...]) -> str:
    return _ROW_JOIN.join(cells)


def _format_warning_list(items: tuple[str, ...] | list[str]) -> str:
    if not items:
        return ""
    return ",".join(items)


def _add_warning(warnings: list[str], code: str) -> None:
    if code not in warnings:
        warnings.append(code)


def _add_reason(reasons: list[str], warnings: list[str], code: str) -> None:
    if code not in reasons:
        reasons.append(code)
    _add_warning(warnings, code)


def _append_unique(existing: tuple[str, ...], code: str) -> tuple[str, ...]:
    if code in existing:
        return existing
    return (*existing, code)
